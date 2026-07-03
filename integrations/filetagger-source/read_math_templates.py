import openpyxl, sys
sys.stdout.reconfigure(encoding='utf-8', errors='replace')

files = [
    r"\\192.168.2.50\brain\15_TEMPLATES\Master_EQ_Chain.xlsx",
    r"\\192.168.2.50\brain\15_TEMPLATES\MASTER_EQUATION_WORKBOOK.xlsx",
    r"\\192.168.2.50\brain\15_TEMPLATES\MATH_TRANSLATION_MASTER_FIXED.xlsx",
    r"\\192.168.2.50\brain\15_TEMPLATES\MATH_TRANSLATION_TABLE_REAL.xlsx",
]

for fpath in files:
    try:
        wb = openpyxl.load_workbook(fpath, data_only=True)
        fname = fpath.split("\\")[-1]
        print(f"\n{'='*70}")
        print(f"FILE: {fname}")
        print(f"SHEETS: {wb.sheetnames}")
        for sname in wb.sheetnames[:5]:
            ws = wb[sname]
            print(f"\n  --- {sname} ({ws.max_row} rows x {ws.max_column} cols) ---")
            for row in ws.iter_rows(max_row=min(15, ws.max_row), values_only=True):
                vals = [str(v)[:55] if v is not None else '' for v in row]
                clean = ' | '.join(v for v in vals if v)
                if clean.strip():
                    print(f"    {clean[:140]}")
        wb.close()
    except Exception as e:
        print(f"\nERROR reading {fpath.split(chr(92))[-1]}: {e}")
